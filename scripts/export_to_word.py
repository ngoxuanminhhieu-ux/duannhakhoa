"""
Script tự động chuyển đổi Báo cáo Tuần 1 thành tệp Microsoft Word (.docx) chuyên nghiệp.
Hỗ trợ định dạng Title, Headings, Bảng biểu đẹp mắt, Chèn hình ảnh OpenCV trực quan và chú thích.

Tác giả: Ngô Xuân Minh Hiếu
Bài tập lớn: Hệ thống quản lý nha khoa tích hợp AI - Tuần 1
"""

import os
import sys

if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


def set_cell_background(cell, fill_hex: str):
    """Đặt màu nền cho cell trong bảng."""
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Đặt lề trong cho cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)


def add_heading_styled(doc, text, level):
    """Thêm Heading với style chuyên nghiệp."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    
    if level == 1:
        run.font.size = Pt(15)
        run.font.color.rgb = RGBColor(15, 44, 89) # Deep Navy
    elif level == 2:
        run.font.size = Pt(13.5)
        run.font.color.rgb = RGBColor(30, 78, 140)
    elif level == 3:
        run.font.size = Pt(12.5)
        run.font.color.rgb = RGBColor(50, 50, 50)
    return p


def create_word_report(output_docx_path: str):
    doc = docx.Document()
    
    # Cấu hình Lề trang (1 inch = 2.54 cm)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base Font style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(40, 40, 40)
    normal_style.paragraph_format.line_spacing = 1.25
    normal_style.paragraph_format.space_after = Pt(6)

    # --- TRANG BÌA / HEADER DỰ ÁN ---
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_sub = title_p.add_run("BÀI TẬP LỚN HỌC PHẦN: THỊ GIÁC MÁY TÍNH & TRÍ TUỆ NHÂN TẠO\n")
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(100, 100, 100)

    run_title = title_p.add_run("BÁO CÁO KẾT QUẢ THỰC HIỆN TUẦN 1\n")
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 44, 89)

    run_subtitle = title_p.add_run("ĐỀ TÀI: HỆ THỐNG QUẢN LÝ NHA KHOA TÍCH HỢP AI DIAGNOSTICS")
    run_subtitle.font.size = Pt(14)
    run_subtitle.font.bold = True
    run_subtitle.font.color.rgb = RGBColor(192, 57, 43)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Đóng khung thông tin sinh viên
    info_table = doc.add_table(rows=4, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False
    
    info_data = [
        ("Sinh viên thực hiện:", "Ngô Xuân Minh Hiếu"),
        ("Nhiệm vụ Tuần 1:", "Khảo sát tập dữ liệu ảnh X-quang nha khoa Panorama & Xác định bài toán AI"),
        ("Công nghệ sử dụng:", "Python, OpenCV, Kaggle, Roboflow, Matplotlib, PyTorch / YOLOv8"),
        ("Thời gian thực hiện:", "Tuần 1 - Học kỳ hiện tại")
    ]
    
    for idx, (label, val) in enumerate(info_data):
        row = info_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.3)
        
        p_l = cell_lbl.paragraphs[0]
        r_l = p_l.add_run(label)
        r_l.bold = True
        r_l.font.color.rgb = RGBColor(15, 44, 89)
        p_l.paragraph_format.space_after = Pt(2)

        p_v = cell_val.paragraphs[0]
        r_v = p_v.add_run(val)
        p_v.paragraph_format.space_after = Pt(2)
        
        set_cell_background(cell_lbl, "F0F4F8")
        set_cell_background(cell_val, "FAFAFA")
        set_cell_margins(cell_lbl, 80, 80, 100, 100)
        set_cell_margins(cell_val, 80, 80, 100, 100)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # --- 1. TỔNG QUAN VỀ ẢNH X-QUANG NHA KHOA PANORAMA ---
    add_heading_styled(doc, "1. TỔNG QUAN VỀ ẢNH X-QUANG NHA KHOA PANORAMA (OPG)", level=1)
    
    add_heading_styled(doc, "1.1 Khái niệm ảnh Panorama (Orthopantomogram - OPG)", level=2)
    p = doc.add_paragraph("Ảnh X-quang Nha khoa Panorama (OPG) là phương pháp chụp quét toàn cảnh hệ thống hai hàm răng, xương hàm trên, xương hàm dưới, khớp thái dương hàm và xoang hàm trên chỉ trên một hình ảnh duy nhất. Đây là công cụ chẩn đoán hình ảnh đầu tay và quy chuẩn trong ngành nha khoa hiện đại.")
    
    add_heading_styled(doc, "1.2 Thách thức kỹ thuật đối với Bài toán Trí tuệ Nhân tạo (AI)", level=2)
    challenges = [
        ("Độ tương phản thấp (Low Contrast): ", "Độ chênh lệch mức xám giữa men răng, ngà răng, tủy răng và các lỗ sâu răng/tổn thương thấu quang cực kỳ nhỏ."),
        ("Nhiễu máy chụp (Speckle Noise & Sensor Artifacts): ", "Ảnh X-quang bị ảnh hưởng bởi nhiễu hạt do bức xạ tia X và máy quét cảm biến digital."),
        ("Tỉ lệ khung hình bất đối xứng (Aspect Ratio): ", "Ảnh Panorama thường rộng khổ 1500 × 800 pixel (tỉ lệ ~ 2:1), không phải hình vuông chuẩn 640 × 640 như đầu vào các mạng CNN/YOLO."),
        ("Nhiễu giải phẫu (Overlapping Structures): ", "Hình ảnh bị chồng lấp bởi bóng của cột sống cổ, xoang hàm và sự chen chúc của các mầm răng.")
    ]
    for title, desc in challenges:
        bp = doc.add_paragraph(style='List Bullet')
        r1 = bp.add_run(title)
        r1.bold = True
        r1.font.color.rgb = RGBColor(15, 44, 89)
        bp.add_run(desc)

    # --- 2. KHẢO SÁT CÁC TẬP DỮ LIỆU ---
    add_heading_styled(doc, "2. KHẢO SÁT CÁC TẬP DỮ LIỆU ẢNH X-QUANG NHA KHOA (DATASETS SURVEY)", level=1)
    doc.add_paragraph("Chúng tôi đã tiến hành nghiên cứu và đánh giá 5 tập dữ liệu ảnh X-quang nha khoa công khai uy tín nhất hiện nay trên Kaggle, Roboflow Universe và cuộc thi MICCAI DENTEX Challenge:")

    # Bảng khảo sát dữ liệu
    table = doc.add_table(rows=6, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    headers = ["Tập Dữ Liệu", "Nguồn / Nền Tảng", "Quy Mô", "Định Dạng Nhãn", "Bài Toán Phù Hợp"]
    hdr_row = table.rows[0]
    for idx, header_text in enumerate(headers):
        cell = hdr_row.cells[idx]
        set_cell_background(cell, "0F2C59")
        set_cell_margins(cell, 120, 120, 100, 100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header_text)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(10.5)

    ds_data = [
        ("DENTEX Benchmark", "MICCAI 2023 Challenge", "~1,000 ảnh Panorama", "Bounding Box + FDI Tooth Numbering", "Phát hiện bệnh lý + Đánh số răng FDI (Khuyên dùng)"),
        ("Tufts Dental Database", "Tufts University", "1,000 ảnh Panorama", "Multi-class Mask / Segmentation", "Phân đoạn răng & bệnh lý chi tiết"),
        ("Dental X-Ray Panoramic", "Roboflow Universe", "500 - 2,000+ ảnh", "YOLO format txt / COCO JSON", "Huấn luyện nhanh mô hình YOLOv8"),
        ("Adult Caries Detection", "Kaggle", "~1,000 ảnh Panorama", "Bounding Box", "Phát hiện sâu răng người lớn"),
        ("Children's Dental", "Kaggle", "~1,500 ảnh Panorama", "Multi-label Bounding Box (14 lớp)", "Phân tích nha khoa trẻ em")
    ]

    for r_idx, row_values in enumerate(ds_data):
        row = table.rows[r_idx + 1]
        bg_color = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, 100, 100, 80, 80)
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.size = Pt(9.5)
            if c_idx == 0:
                run.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # --- 3. XÁC ĐỊNH BÀI TOÁN AI ---
    add_heading_styled(doc, "3. XÁC ĐỊNH & MÔ HÌNH HÓA CÁC BÀI TOÁN AI TRONG NHA KHOA", level=1)
    
    add_heading_styled(doc, "3.1 Bài toán 0 (Nền tảng): Định danh & Đánh số răng chuẩn Quốc tế FDI", level=2)
    doc.add_paragraph("Mục tiêu định vị 32 chiếc răng theo hệ thống đánh số FDI (11–18: Hàm trên phải; 21–28: Hàm trên trái; 31–38: Hàm dưới trái; 41–48: Hàm dưới phải). Đây là bài toán bản đồ cốt lõi giúp kết quả chẩn đoán tự động đồng bộ trực tiếp vào Sơ đồ răng (Dental Chart) trong phần mềm Quản lý Nha khoa.")

    add_heading_styled(doc, "3.2 Bài toán 1: Chẩn đoán & Phân cấp Sâu răng (Dental Caries Detection)", level=2)
    doc.add_paragraph("Mục tiêu phát hiện vị trí sâu răng và phân cấp độ nghiêm trọng từ sâu men răng (C1-C2), sâu ngà răng (C3) đến sâu sát tủy (C4). Mô hình đề xuất: YOLOv8 / Segment Anything (SAM).")

    add_heading_styled(doc, "3.3 Bài toán 2: Chẩn đoán Răng khôn mọc lệch (Impacted Wisdom Teeth)", level=2)
    doc.add_paragraph("Nhận diện các răng số 18, 28, 38, 48 và phân loại kiểu mọc lệch theo tiêu chuẩn Winter: Mọc thẳng (Vertical), Nghiêng gần (Mesioangular), Nghiêng xa (Distoangular), Nằm ngang (Horizontal), Mọc ngầm hoàn toàn trong xương.")

    add_heading_styled(doc, "3.4 Bài toán 3: Chẩn đoán Viêm cuống răng & Viêm nha chu (Periapical Lesions & Periodontitis)", level=2)
    doc.add_paragraph("Phát hiện các vùng thấu quang tối màu ở chóp chân răng (áp-xe, nang chân răng) và tình trạng sụt giảm chiều cao xương ổ răng.")

    # --- 4. THỬ NGHIỆM TIỀN XỬ LÝ OPENCV ---
    add_heading_styled(doc, "4. KẾT QUẢ TRIỂN KHAI TIỀN XỬ LÝ ẢNH X-QUANG VỚI OPENCV", level=1)
    doc.add_paragraph("Sinh viên đã xây dựng module Python OpenCV (src/preprocessing.py) ứng dụng thuật toán CLAHE (Contrast Limited Adaptive Histogram Equalization) kết hợp Fast NLM Denoising.")

    # Thêm bảng kết quả đo chỉ số tương phản
    metrics_table = doc.add_table(rows=4, cols=4)
    metrics_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    m_headers = ["Trạng Thái Ảnh", "RMS Contrast", "Độ Lệch Chuẩn (Std)", "Đánh Giá Thực Tế"]
    m_hdr_row = metrics_table.rows[0]
    for idx, h_text in enumerate(m_headers):
        cell = m_hdr_row.cells[idx]
        set_cell_background(cell, "0F2C59")
        set_cell_margins(cell, 100, 100, 100, 100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(10)

    m_data = [
        ("Ảnh X-quang Gốc", "0.1071", "27.31", "Ảnh mờ, tương phản kém, khó nhận biết vết sâu nhỏ"),
        ("Sau xử lý CLAHE", "0.1698", "43.31", "Độ tương phản tăng 58.5%, ranh giới răng hiện rõ"),
        ("Sau xử lý Denoising", "0.1696", "43.24", "Giữ độ tương phản cao, khử mịn nhiễu hạt X-quang")
    ]

    for r_idx, row_values in enumerate(m_data):
        row = metrics_table.rows[r_idx + 1]
        bg_color = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, 80, 80, 80, 80)
            p = cell.paragraphs[0]
            if c_idx in [1, 2]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val)
            run.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # --- CHÈN HÌNH ẢNH MINH HỌA VÀO FILE WORD ---
    img1_path = "data/samples/preprocessing_comparison.png"
    if os.path.exists(img1_path):
        add_heading_styled(doc, "4.1 Hình ảnh So sánh Kỹ thuật Tiền xử lý OpenCV", level=2)
        p_img1 = doc.add_paragraph()
        p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img1.add_run().add_picture(img1_path, width=Inches(6.2))
        
        cap1 = doc.add_paragraph()
        cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap1 = cap1.add_run("Hình 1: So sánh biểu đồ Histogram và chất lượng ảnh trước và sau khi xử lý CLAHE & Denoising.")
        r_cap1.font.italic = True
        r_cap1.font.size = Pt(9.5)
        r_cap1.font.color.rgb = RGBColor(100, 100, 100)

    img2_path = "data/samples/annotated_panoramic_demo.png"
    if os.path.exists(img2_path):
        add_heading_styled(doc, "4.2 Hình ảnh Minh họa Gán nhãn Bệnh lý & Đánh số răng FDI", level=2)
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.add_run().add_picture(img2_path, width=Inches(6.2))
        
        cap2 = doc.add_paragraph()
        cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap2 = cap2.add_run("Hình 2: Mô phỏng phát hiện Bounding Box các bệnh lý Sâu răng, Răng khôn mọc lệch kết hợp mã răng FDI.")
        r_cap2.font.italic = True
        r_cap2.font.size = Pt(9.5)
        r_cap2.font.color.rgb = RGBColor(100, 100, 100)

    # --- 5. CẤU TRÚC MÃ NGUỒN & HƯỚNG DẪN ---
    add_heading_styled(doc, "5. CẤU TRÚC MÃ NGUỒN VÀ HƯỚNG DẪN THỰC THI", level=1)
    doc.add_paragraph("Mã nguồn dự án đã được thiết kế hoàn chỉnh trong thư mục c:\\Duannhakhoa bao gồm:")
    
    files_list = [
        ("BAO_CAO_TUAN_1.md / BAO_CAO_TUAN_1.docx: ", "Báo cáo tổng hợp kết quả Tuần 1."),
        ("src/preprocessing.py: ", "Module OpenCV tiền xử lý ảnh X-quang (CLAHE, Denoising, Letterbox)."),
        ("src/eda_visualization.py: ", "Module hiển thị biểu đồ và gán nhãn bệnh lý."),
        ("src/dataset_downloader.py: ", "Module hỗ trợ tải dữ liệu từ Kaggle & Roboflow API."),
        ("scripts/generate_sample_demo.py: ", "Script chạy tự động toàn bộ quy trình tiền xử lý và sinh ảnh báo cáo.")
    ]
    for title, desc in files_list:
        bp = doc.add_paragraph(style='List Bullet')
        r1 = bp.add_run(title)
        r1.bold = True
        bp.add_run(desc)

    # --- 6. KẾ HOẠCH TUẦN 2 ---
    add_heading_styled(doc, "6. KẾ HOẠCH TRIỂN KHAI CHO TUẦN 2", level=1)
    plans = [
        "Tải và tiền xử lý bộ dữ liệu DENTEX / Roboflow Dental X-Ray Dataset chuẩn định dạng YOLOv8.",
        "Thiết lập bài toán và huấn luyện mô hình YOLOv8s cho bài toán Chẩn đoán Sâu răng và Răng khôn mọc lệch.",
        "Đánh giá độ chính xác mô hình thông qua các chỉ số mAP@0.5, Precision, Recall và Confusion Matrix.",
        "Xây dựng API Python (FastAPI/Flask) sẵn sàng kết nối với giao diện phần mềm Quản lý Nha khoa."
    ]
    for idx, plan_item in enumerate(plans, 1):
        p = doc.add_paragraph()
        r = p.add_run(f"6.{idx} {plan_item}")

    # Chữ ký
    doc.add_paragraph().paragraph_format.space_after = Pt(20)
    sig_p = doc.add_paragraph()
    sig_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    sig_run1 = sig_p.add_run("Sinh viên thực hiện\n\n\n\n")
    sig_run1.font.bold = True
    sig_run2 = sig_p.add_run("Ngô Xuân Minh Hiếu")
    sig_run2.font.bold = True
    sig_run2.font.size = Pt(13)

    # Lưu file Word
    doc.save(output_docx_path)
    print(f"=== Đã tạo thành công tệp Microsoft Word (.docx) tại: {output_docx_path} ===")


if __name__ == "__main__":
    output_path = "BAO_CAO_TUAN_1.docx"
    create_word_report(output_path)
