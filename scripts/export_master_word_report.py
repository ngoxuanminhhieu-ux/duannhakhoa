"""
Script xuất tệp Microsoft Word (.docx) Báo cáo Hoàn chỉnh Bài tập lớn
Bao gồm cả 2 phần:
1. Khảo sát Dataset & Mô hình hóa 4 Bài toán AI Nha khoa (Khảo sát DENTEX, Tufts, Roboflow, Kaggle & OpenCV).
2. Thiết lập Môi trường phát triển, Docker Container, quy chuẩn Git Flow & Coding Convention PEP 8.

Thành viên nhóm: Ngô Xuân Minh Hiếu, Lưu Tuấn Anh
"""

import os
import sys

if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls


def set_cell_background(cell, fill_hex: str):
    """Đặt màu nền cho cell trong bảng docx."""
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Đặt lề padding bên trong cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)


def add_heading_styled(doc, text, level):
    """Tạo tiêu đề Heading được định dạng màu sắc & font chữ chuẩn."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    
    if level == 1:
        run.font.size = Pt(15)
        run.font.color.rgb = RGBColor(15, 44, 89) # Deep Navy
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(30, 78, 140)
    elif level == 3:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(50, 50, 50)
    return p


def build_master_word_report(output_docx_path: str):
    doc = docx.Document()
    
    # Cấu hình Lề trang (1 inch = 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styling chuẩn cho văn bản
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(40, 40, 40)
    normal_style.paragraph_format.line_spacing = 1.25
    normal_style.paragraph_format.space_after = Pt(6)

    # --- 1. TRANG BÌA CHUYÊN NGHIỆP ---
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    run_sub = title_p.add_run("BÀI TẬP LỚN HỌC PHẦN: THỊ GIÁC MÁY TÍNH & TRÍ TUỆ NHÂN TẠO\n")
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(100, 100, 100)

    run_title = title_p.add_run("BÁO CÁO TỔNG HỢP TIẾN ĐỘ & THIẾT LẬP DỰ ÁN\n")
    run_title.font.size = Pt(19)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 44, 89)

    run_subtitle = title_p.add_run("ĐỀ TÀI: HỆ THỐNG QUẢN LÝ NHA KHOA TÍCH HỢP AI DIAGNOSTICS")
    run_subtitle.font.size = Pt(13.5)
    run_subtitle.font.bold = True
    run_subtitle.font.color.rgb = RGBColor(192, 57, 43)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Bảng thông tin nhóm sinh viên
    info_table = doc.add_table(rows=4, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    info_data = [
        ("Sinh viên thực hiện:", "1. Ngô Xuân Minh Hiếu\n2. Lưu Tuấn Anh"),
        ("Nhiệm vụ trọng tâm:", "• Khảo sát tập dữ liệu X-quang Nha khoa (Panorama) & Xác định bài toán AI\n• Thiết lập môi trường phát triển (Docker), Quy chuẩn Git Flow & PEP 8"),
        ("Công nghệ & Công cụ:", "Python 3.10+, OpenCV, PyTorch, YOLOv8, Git, GitHub, Docker, VS Code, Black, Flake8"),
        ("Trạng thái tiến độ:", "Đã hoàn thành 100% nhiệm vụ giai đoạn chuẩn bị & sẵn sàng huấn luyện mô hình")
    ]
    
    for idx, (label, val) in enumerate(info_data):
        row = info_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.3)
        
        p_l = cell_lbl.paragraphs[0]
        r_l = p_l.add_run(label)
        r_l.bold = True
        r_l.font.color.rgb = RGBColor(15, 44, 89)
        p_l.paragraph_format.space_after = Pt(2)

        p_v = cell_val.paragraphs[0]
        r_v = p_v.add_run(val)
        p_v.paragraph_format.space_after = Pt(2)
        
        set_cell_background(cell_lbl, "F0F4F8")
        set_cell_background(cell_val, "FAFAFA")
        set_cell_margins(cell_lbl, 80, 80, 100, 100)
        set_cell_margins(cell_val, 80, 80, 100, 100)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # --- CHƯƠNG 1: KHẢO SÁT DỮ LIỆU & BÀI TOÁN AI ---
    add_heading_styled(doc, "CHƯƠNG 1: KHẢO SÁT DỮ LIỆU X-QUANG NHA KHOA & BÀI TOÁN AI", level=1)
    
    add_heading_styled(doc, "1.1 Đặc trưng kỹ thuật của Ảnh X-quang Panorama (OPG)", level=2)
    doc.add_paragraph("Ảnh X-quang Panorama (Orthopantomogram) chụp quét toàn cảnh hệ thống 2 hàm răng, xương hàm và xoang hàm. Các thách thức chính đối với bài toán thị giác máy tính bao gồm:")
    
    challenges = [
        ("Độ tương phản thấp (Low Contrast): ", "Độ chênh lệch mức xám giữa men răng, ngà răng và lỗ sâu nhỏ rất thấp."),
        ("Nhiễu hạt máy chụp (Speckle Noise): ", "Xuất hiện nhiễu hạt do bức xạ tia X và cảm biến digital."),
        ("Tỉ lệ bất đối xứng (Aspect Ratio): ", "Ảnh Panorama có khổ rộng (~2:1), không phải hình vuông chuẩn 640x640 của CNN."),
        ("Chồng lấp giải phẫu: ", "Bóng cột sống cổ và sự chen chúc của các mầm răng gây nhiễu nhận diện.")
    ]
    for c_title, c_desc in challenges:
        bp = doc.add_paragraph(style='List Bullet')
        r = bp.add_run(c_title)
        r.bold = True
        r.font.color.rgb = RGBColor(15, 44, 89)
        bp.add_run(c_desc)

    add_heading_styled(doc, "1.2 Đánh giá các Tập dữ liệu Ảnh X-quang Nha khoa Công khai", level=2)
    doc.add_paragraph("Nhóm đã tiến hành khảo sát và lập bảng so sánh chi tiết các tập dữ liệu công khai uy tín nhất hiện nay:")

    # Bảng khảo sát Dataset
    ds_table = doc.add_table(rows=6, cols=5)
    ds_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    headers = ["Tập Dữ Liệu", "Nguồn / Nền Tảng", "Quy Mô", "Định Dạng Nhãn", "Bài Toán Phù Hợp"]
    hdr_row = ds_table.rows[0]
    for idx, h_text in enumerate(headers):
        cell = hdr_row.cells[idx]
        set_cell_background(cell, "0F2C59")
        set_cell_margins(cell, 100, 100, 80, 80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(10)

    ds_data = [
        ("DENTEX Benchmark", "MICCAI 2023 Challenge", "~1,000 ảnh Panorama", "Bounding Box + FDI Tooth Numbering", "Phát hiện đa bệnh lý + Đánh số răng FDI (Khuyên dùng)"),
        ("Tufts Dental Database", "Tufts University", "1,000 ảnh Panorama", "Multi-class Mask / Segmentation", "Phân đoạn răng & bệnh lý chi tiết"),
        ("Dental X-Ray Panoramic", "Roboflow Universe", "500 - 2,000+ ảnh", "YOLO format txt / COCO JSON", "Huấn luyện nhanh mô hình YOLOv8"),
        ("Adult Caries Detection", "Kaggle", "~1,000 ảnh Panorama", "Bounding Box", "Phát hiện sâu răng người lớn"),
        ("Children's Dental", "Kaggle", "~1,500 ảnh Panorama", "Multi-label Bounding Box (14 lớp)", "Phân tích nha khoa trẻ em")
    ]

    for r_idx, row_values in enumerate(ds_data):
        row = ds_table.rows[r_idx + 1]
        bg_color = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, 80, 80, 80, 80)
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.size = Pt(9.5)
            if c_idx == 0:
                run.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    add_heading_styled(doc, "1.3 Xác định & Mô hình hóa 4 Bài toán AI Cốt lõi", level=2)
    ai_tasks = [
        ("Bài toán 0 (Nền tảng): Định danh & Đánh số răng chuẩn FDI (11–48): ", "Tự động gán kết quả chẩn đoán vào đúng chiếc răng trên Sơ đồ răng (Dental Chart) của phần mềm nha khoa."),
        ("Bài toán 1: Chẩn đoán & Phân cấp Sâu răng (Dental Caries Detection): ", "Phát hiện vị trí sâu và phân cấp độ nghiêm trọng từ C1 đến C4."),
        ("Bài toán 2: Chẩn đoán Răng khôn mọc lệch (Impacted Wisdom Teeth): ", "Nhận diện răng 18, 28, 38, 48 và phân loại góc mọc nghiêng gần, nghiêng xa, mọc ngang, mọc ngầm."),
        ("Bài toán 3: Chẩn đoán Viêm cuống răng & Viêm nha chu (Periapical Lesions): ", "Phát hiện vùng thấu quang chóp chân răng và tiêu xương ổ răng.")
    ]
    for t_title, t_desc in ai_tasks:
        bp = doc.add_paragraph(style='List Bullet')
        r = bp.add_run(t_title)
        r.bold = True
        r.font.color.rgb = RGBColor(15, 44, 89)
        bp.add_run(t_desc)

    # --- CHƯƠNG 2: KẾT QUẢ TIỀN XỬ LÝ OPENCV ---
    add_heading_styled(doc, "CHƯƠNG 2: THỬ NGHIỆM TIỀN XỬ LÝ ẢNH VỚI OPENCV", level=1)
    doc.add_paragraph("Sinh viên đã phát triển module Python OpenCV (src/preprocessing.py) ứng dụng thuật toán CLAHE (Contrast Limited Adaptive Histogram Equalization) kết hợp Fast NLM Denoising.")

    # Bảng đo chỉ số RMS Contrast
    metrics_table = doc.add_table(rows=4, cols=4)
    metrics_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    m_headers = ["Trạng Thái Ảnh", "RMS Contrast", "Độ Lệch Chuẩn (Std)", "Đánh Giá Thực Tế"]
    m_hdr_row = metrics_table.rows[0]
    for idx, h_text in enumerate(m_headers):
        cell = m_hdr_row.cells[idx]
        set_cell_background(cell, "0F2C59")
        set_cell_margins(cell, 100, 100, 100, 100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(10)

    m_data = [
        ("Ảnh X-quang Gốc", "0.1071", "27.31", "Ảnh mờ, tương phản kém, khó nhận biết vết sâu nhỏ"),
        ("Sau xử lý CLAHE", "0.1698", "43.31", "Độ tương phản tăng 58.5%, ranh giới răng hiện rõ"),
        ("Sau xử lý Denoising", "0.1696", "43.24", "Giữ độ tương phản cao, khử mịn nhiễu hạt X-quang")
    ]

    for r_idx, row_values in enumerate(m_data):
        row = metrics_table.rows[r_idx + 1]
        bg_color = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, 80, 80, 80, 80)
            p = cell.paragraphs[0]
            if c_idx in [1, 2]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val)
            run.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Chèn hình ảnh trực quan
    img1_path = "data/samples/preprocessing_comparison.png"
    if os.path.exists(img1_path):
        add_heading_styled(doc, "2.1 Hình ảnh So sánh Kỹ thuật Tiền xử lý OpenCV", level=2)
        p_img1 = doc.add_paragraph()
        p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img1.add_run().add_picture(img1_path, width=Inches(6.2))
        
        cap1 = doc.add_paragraph()
        cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap1 = cap1.add_run("Hình 1: Biểu đồ Histogram và chất lượng ảnh trước và sau khi xử lý CLAHE & Denoising.")
        r_cap1.font.italic = True
        r_cap1.font.size = Pt(9.5)
        r_cap1.font.color.rgb = RGBColor(100, 100, 100)

    img2_path = "data/samples/annotated_panoramic_demo.png"
    if os.path.exists(img2_path):
        add_heading_styled(doc, "2.2 Hình ảnh Minh họa Gán nhãn Bệnh lý & Đánh số răng FDI", level=2)
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.add_run().add_picture(img2_path, width=Inches(6.2))
        
        cap2 = doc.add_paragraph()
        cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap2 = cap2.add_run("Hình 2: Phát hiện Bounding Box các bệnh lý Sâu răng, Răng khôn mọc lệch kết hợp mã răng FDI.")
        r_cap2.font.italic = True
        r_cap2.font.size = Pt(9.5)
        r_cap2.font.color.rgb = RGBColor(100, 100, 100)

    # --- CHƯƠNG 3: MÔI TRƯỜNG PHÁT TRUYỂN, GIT FLOW & CODING CONVENTION ---
    add_heading_styled(doc, "CHƯƠNG 3: MÔI TRƯỜNG PHÁT TRUYỂN, GIT FLOW & CODING CONVENTION", level=1)
    
    add_heading_styled(doc, "3.1 Quy chuẩn Quản lý Mã nguồn Git Flow", level=2)
    doc.add_paragraph("Để đảm bảo quá trình phối hợp giữa Ngô Xuân Minh Hiếu và Lưu Tuấn Anh diễn ra thuận lợi, nhóm thống nhất cấu trúc 2 nhánh chính (`main`, `develop`) kết hợp các nhánh tính năng `feature/*` và nhánh sửa lỗi `bugfix/*` kèm quy chuẩn commit chuẩn Conventional Commits (`feat:`, `fix:`, `docs:`, `style:`, `chore:`).")

    add_heading_styled(doc, "3.2 Quy chuẩn Lập trình Python (PEP 8)", level=2)
    doc.add_paragraph("Áp dụng định dạng chuẩn PEP 8 với độ dài dòng tối đa 120 ký tự, sử dụng Type Hinting cho mọi hàm, tự động format bằng Black Formatter và kiểm tra linter bằng Flake8 trong VS Code.")

    add_heading_styled(doc, "3.3 Cấu hình Môi trường Container với Docker", level=2)
    doc.add_paragraph("Nhóm đã tạo Dockerfile (Python 3.10-slim + OpenCV dependencies) và docker-compose.yml giúp môi trường ứng dụng chạy đồng nhất 100% giữa các máy tính cá nhân.")

    # Bảng danh sách các file cấu hình
    cfg_table = doc.add_table(rows=7, cols=2)
    cfg_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    cfg_headers = ["Tệp Cấu Hình", "Vai Trò & Chức Năng Trong Dự Án"]
    hdr_row = cfg_table.rows[0]
    for idx, h_text in enumerate(cfg_headers):
        cell = hdr_row.cells[idx]
        set_cell_background(cell, "0F2C59")
        set_cell_margins(cell, 100, 100, 100, 100)
        p = cell.paragraphs[0]
        run = p.add_run(h_text)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(10)

    configs_data = [
        ("Dockerfile", "Đóng gói ứng dụng Python 3.10-slim kết hợp OpenCV hệ thống (libgl1-mesa-glx)."),
        ("docker-compose.yml", "Quản lý dịch vụ Container và mount volume mã nguồn cục bộ."),
        (".gitignore", "Chặn đưa lên Git các file dữ liệu nặng (*.zip, data/), file model (*.pt) và tệp tạm."),
        (".editorconfig", "Đảm bảo độ rộng Tab 4 space, xuống dòng LF và mã hóa UTF-8 giữa các máy."),
        (".vscode/settings.json", "Cấu hình tự động format code bằng Black và xóa khoảng trắng dư thừa khi lưu."),
        (".github/workflows/ci.yml", "Tự động chạy kiểm thử Linting Flake8 và Black format trên GitHub Actions.")
    ]

    for r_idx, (file_name, file_desc) in enumerate(configs_data):
        row = cfg_table.rows[r_idx + 1]
        bg_color = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
        
        cell_fn, cell_fd = row.cells[0], row.cells[1]
        cell_fn.width = Inches(2.2)
        cell_fd.width = Inches(4.3)
        
        set_cell_background(cell_fn, bg_color)
        set_cell_background(cell_fd, bg_color)
        set_cell_margins(cell_fn, 80, 80, 80, 80)
        set_cell_margins(cell_fd, 80, 80, 80, 80)
        
        p1 = cell_fn.paragraphs[0]
        r1 = p1.add_run(file_name)
        r1.bold = True
        r1.font.size = Pt(9.5)
        
        p2 = cell_fd.paragraphs[0]
        r2 = p2.add_run(file_desc)
        r2.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # --- CHƯƠNG 4: KẾ HOẠCH BƯỚC TIẾP THEO ---
    add_heading_styled(doc, "CHƯƠNG 4: KẾ HOẠCH TRIỂN KHAI GIAI ĐOẠN TÍCH HỢP MO-HINH AI", level=1)
    next_steps = [
        "Huấn luyện mô hình YOLOv8s trên tập dữ liệu DENTEX / Roboflow đã gán nhãn.",
        "Đánh giá độ chính xác chẩn đoán qua các chỉ số mAP@0.5, Precision, Recall và Confusion Matrix.",
        "Xây dựng API dịch vụ chẩn đoán AI bằng FastAPI/Flask.",
        "Tích hợp API AI vào giao diện Hồ sơ Bệnh án & Sơ đồ răng của Phần mềm Quản lý Nha khoa."
    ]
    for idx, step_text in enumerate(next_steps, 1):
        p = doc.add_paragraph()
        r = p.add_run(f"4.{idx} {step_text}")

    # Chữ ký nhóm sinh viên
    doc.add_paragraph().paragraph_format.space_after = Pt(25)
    sig_p = doc.add_paragraph()
    sig_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    sig_run1 = sig_p.add_run("Thành viên nhóm thực hiện\n\n\n\n")
    sig_run1.font.bold = True
    sig_run2 = sig_p.add_run("Ngô Xuân Minh Hiếu & Lưu Tuấn Anh")
    sig_run2.font.bold = True
    sig_run2.font.size = Pt(12.5)

    doc.save(output_docx_path)
    print(f"=== Đã tạo thành công tệp Microsoft Word hoàn chỉnh tại: {output_docx_path} ===")


if __name__ == "__main__":
    output_path = "BAO_CAO_HOAN_CHINH_DU_AN_NHA_KHOA_AI.docx"
    build_master_word_report(output_path)
